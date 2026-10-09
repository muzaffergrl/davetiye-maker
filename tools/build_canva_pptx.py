"""Dusty Blue video master'ını 6 sayfalık katmanlı PPTX olarak kurar (Canva'ya yüklemek için).

Her görsel ayrı resim katmanı, her metin ayrı metin kutusu olur; Canva içe aktarınca
alıcı metinleri ve fotoğrafları değiştirebilir. Animasyonlar PPTX'ten Canva'ya geçmez
(Canva Help Center, 2026-10-09); bunlar Canva arayüzünde eklenir.

Kullanım: python3 tools/build_canva_pptx.py <çıktı.pptx>
Koordinatlar px (1080 × 1920); 1 px = 9525 EMU, font: pt = px × 0.75.
"""
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parent.parent
COL = ROOT / "collections/02_Dusty_Blue"
L = COL / "assets/layers"
C = json.loads((COL / "content.demo.json").read_text())

PX = 9525
INK, BLUE_DEEP, GOLD, IVORY = "2B3A4E", "4E6782", "B8996A", "FBF7EF"
SCRIPT, SERIF, SANS = "Pinyon Script", "Cormorant Garamond", "Montserrat"


def px(v):
    return Emu(int(round(v * PX)))


def img(slide, name, left, top, w, h, rot=0):
    pic = slide.shapes.add_picture(str(L / name), px(left), px(top), px(w), px(h))
    pic.name = name.rsplit(".", 1)[0]
    if rot:
        pic.rotation = rot
    return pic


def text(slide, s, left, top, w, h, font=SERIF, size=40, color=INK, align="c",
         italic=False, bold=False, spacing=0, line=1.15, name=None):
    tb = slide.shapes.add_textbox(px(left), px(top), px(w), px(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    for i, part in enumerate(s.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"c": PP_ALIGN.CENTER, "l": PP_ALIGN.LEFT, "r": PP_ALIGN.RIGHT}[align]
        p.line_spacing = line
        # Pinyon Script'in "&" glifi "v" gibi okunuyor: ampersand Cormorant italik
        chunks = [part]
        if font == SCRIPT and "&" in part:
            a, b = part.split("&", 1)
            chunks = [a, "&", b]
        for chunk in chunks:
            if not chunk:
                continue
            r = p.add_run()
            r.text = chunk
            f = r.font
            amp = chunk == "&" and font == SCRIPT
            f.name = SERIF if amp else font
            f.size = Pt(size * (0.62 if amp else 0.75))
            f.italic = True if amp else italic
            f.bold = bold
            f.color.rgb = RGBColor.from_string(GOLD if amp and color == INK else color)
            if spacing:
                r._r.get_or_add_rPr().set("spc", str(int(spacing * 75)))
    return tb


def line(slide, x1, y, x2, color=GOLD, weight=2):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(x1), px(y), px(x2 - x1), px(weight))
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor.from_string(color)
    shp.line.fill.background()
    return shp


def base(prs, decor_bottom=True):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    img(s, "bg_mist.jpg", 0, 0, 1080, 1920)
    img(s, "botanical_tl.png", 0, 0, 400, 400)
    if decor_bottom:
        img(s, "botanical_br.png", 670, 1510, 460, 460)
    return s


def card(s):
    img(s, "invitation_card_blank.png", 20, 140, 1040, 1640)
    img(s, "card_botanical_tl.png", 94, 214, 250, 250)
    img(s, "card_botanical_br.png", 736, 1456, 250, 250)


def seal(s, left, top, size, mono_px):
    img(s, "wax_seal.png", left, top, size, size)
    c = size / 2
    text(s, C["initials"], left, top + c - mono_px * 0.7, size, mono_px * 1.3,
         size=mono_px, color=INK, name="seal_monogram")


def fmt_time(hhmm):
    h, m = map(int, hhmm.split(":"))
    return f"{(h - 1) % 12 + 1}:{m:02d} {'AM' if h < 12 else 'PM'}"


def main(out):
    import datetime as dt
    d = dt.date.fromisoformat(C["date"])
    rsvp = dt.date.fromisoformat(C["rsvp"]["deadline"])
    p1, p2 = C["partners"]
    v = C["venue"]

    prs = Presentation()
    prs.slide_width, prs.slide_height = px(1080), px(1920)

    # 1 — kapalı zarf
    s = base(prs)
    text(s, "A LOVE LETTER FROM", 140, 296, 800, 40, size=28, color=BLUE_DEEP, spacing=3)
    text(s, C["name_display"], 40, 340, 1000, 170, font=SCRIPT, size=132)
    img(s, "envelope_back.png", 30, 690, 1020, 740)
    img(s, "envelope_pocket.png", 90, 750, 900, 620)
    img(s, "envelope_flap_closed.png", 86, 746, 908, 408)
    seal(s, 416, 1008, 248, 64)
    text(s, "a letter for you, to open with love", 140, 1452, 800, 60, size=44,
         italic=True, color=BLUE_DEEP)

    # 2 — zarf açık, kart yükselmiş (cep önde)
    s = base(prs)
    img(s, "envelope_back.png", 30, 690, 1020, 740)
    img(s, "envelope_liner_open.png", 96, 384, 888, 380)
    img(s, "card_small_blank.png", 130, 150, 820, 600)
    img(s, "photo_placeholder_wide.jpg", 200, 220, 680, 310)
    text(s, "We're getting married", 160, 528, 760, 110, font=SCRIPT, size=80)
    text(s, d.strftime("%B %-d, %Y").upper(), 160, 652, 760, 40, size=26,
         color=BLUE_DEEP, spacing=2)
    img(s, "envelope_pocket.png", 90, 750, 900, 620)
    text(s, "and you are warmly invited", 140, 1454, 800, 60, size=44,
         italic=True, color=BLUE_DEEP)

    # 3 — Our Story, üç polaroid
    s = base(prs)
    text(s, "Our Story", 140, 170, 800, 170, font=SCRIPT, size=130)
    text(s, "A FEW OF OUR FAVORITE MOMENTS", 140, 366, 800, 40, size=26,
         color=BLUE_DEEP, spacing=3)
    pols = [(20, 420, "photo_placeholder_1.jpg", "where it began"),
            (514, 560, "photo_placeholder_2.jpg", "the question"),
            (260, 1060, "photo_placeholder_3.jpg", "forever starts now")]
    for left, top, ph, cap in pols:
        img(s, "polaroid_frame.png", left, top, 540, 630)
        img(s, ph, left + 72, top + 70, 396, 416)
        text(s, cap, left + 70, top + 494, 400, 70, font=SCRIPT, size=46)

    # 4 — davet kartı
    s = base(prs)
    card(s)
    text(s, C["host_line"].upper(), 140, 434, 800, 40, size=26, color=BLUE_DEEP, spacing=2)
    text(s, p1["full_name"], 110, 505, 860, 150, font=SCRIPT, size=112)
    text(s, "&", 440, 672, 200, 100, size=76, italic=True, color=GOLD)
    text(s, p2["full_name"], 110, 775, 860, 150, font=SCRIPT, size=112)
    text(s, "request the pleasure of your company\nat the celebration of their marriage",
         150, 980, 780, 120, size=40, italic=True, line=1.2)
    text(s, d.strftime("%B").upper(), 340, 1156, 400, 36, size=24, color=BLUE_DEEP, spacing=2)
    line(s, 160, 1230, 434)
    line(s, 160, 1300, 434)
    text(s, d.strftime("%A").upper(), 160, 1246, 274, 40, size=24, color=BLUE_DEEP, spacing=2)
    text(s, str(d.day), 440, 1196, 200, 136, size=110)
    line(s, 646, 1230, 920)
    line(s, 646, 1300, 920)
    text(s, "AT " + fmt_time(C["ceremony_time"]), 646, 1246, 274, 40, size=24,
         color=BLUE_DEEP, spacing=2)
    text(s, str(d.year), 340, 1364, 400, 36, size=26, color=BLUE_DEEP, spacing=2)
    text(s, v["name"], 140, 1440, 800, 70, size=54)
    text(s, f'{v["city"]}, {v["region"]}', 140, 1514, 800, 54, size=40, italic=True,
         color=BLUE_DEEP)

    # 5 — program
    s = base(prs)
    card(s)
    text(s, "The Celebration", 110, 504, 860, 160, font=SCRIPT, size=120)
    line(s, 330, 700, 750)
    for i, ev in enumerate(C["schedule"]):
        y = 760 + i * 190
        text(s, fmt_time(ev["time"]), 200, y + 20, 220, 50, size=28, color=BLUE_DEEP,
             align="l", spacing=1)
        text(s, ev["label"], 450, y, 460, 70, size=50, align="l")
        text(s, ev["location"], 450, y + 68, 460, 50, size=36, italic=True,
             color=BLUE_DEEP, align="l")
        if i < len(C["schedule"]) - 1:
            line(s, 200, y + 174, 880, color="D9C49E")
    text(s, "ATTIRE", 340, 1346, 400, 36, size=24, color=BLUE_DEEP, spacing=3)
    text(s, C["dress_code"].capitalize(), 140, 1388, 800, 64, size=50, italic=True)

    # 6 — RSVP
    s = base(prs)
    card(s)
    text(s, "Kindly Reply", 110, 480, 860, 160, font=SCRIPT, size=120)
    text(s, "by " + rsvp.strftime("%B %-d, %Y"), 140, 662, 800, 64, size=50,
         italic=True, bold=True)
    btn = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(300), px(800), px(480), px(100))
    btn.adjustments[0] = 0.5
    btn.fill.solid()
    btn.fill.fore_color.rgb = RGBColor.from_string(BLUE_DEEP)
    btn.line.fill.background()
    btn.name = "rsvp_button"
    text(s, "RSVP ONLINE", 300, 800, 480, 100, font=SANS, size=28, color=IVORY,
         spacing=3, name="rsvp_button_label")
    text(s, "Details, travel & registry\non our wedding website", 140, 940, 800, 110,
         size=38, italic=True, color=BLUE_DEEP, line=1.2)
    seal(s, 465, 1118, 150, 38)
    text(s, "We can't wait to\ncelebrate with you", 140, 1300, 800, 200, font=SCRIPT,
         size=68, line=1.15)

    prs.save(out)
    print("saved", out, len(prs.slides), "slides")


if __name__ == "__main__":
    main(sys.argv[1])
